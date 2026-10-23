"""Parse RFC 822 emails and extract text from common business attachments."""

from __future__ import annotations

import csv
import hashlib
import io
from datetime import date, datetime
from email import policy
from email.message import EmailMessage, Message
from email.parser import BytesParser
from email.utils import getaddresses, parseaddr
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from .models import AttachmentRecord, InboundEmail


# These limits prevent an innocent-looking email from consuming unbounded RAM
# or model context. They are intentionally conservative for a learning project.
MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024
MAX_TOTAL_ATTACHMENT_BYTES = 25 * 1024 * 1024
MAX_ATTACHMENTS = 20
MAX_EXTRACTED_CHARACTERS = 200_000
MAX_SPREADSHEET_ROWS = 20_000
MAX_PDF_PAGES = 100


class _HTMLTextExtractor(HTMLParser):
    """Tiny, dependency-free HTML-to-text fallback for email bodies."""

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        text = data.strip()
        if text:
            self.parts.append(text)

    def text(self) -> str:
        return "\n".join(self.parts)


def _addresses(values: Iterable[str]) -> list[str]:
    """Normalize one or more RFC-formatted address headers."""

    return [address.lower() for _, address in getaddresses(values) if address]


def _decode_text(payload: bytes) -> str:
    """Decode user-supplied text with a practical Windows-email fallback."""

    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return payload.decode("utf-8", errors="replace")


def _body_text(message: EmailMessage) -> str:
    """Prefer plain text, falling back to a stripped HTML alternative."""

    plain_parts: list[str] = []
    html_parts: list[str] = []
    parts: Iterable[Message] = message.walk() if message.is_multipart() else [message]
    for part in parts:
        if part.get_content_disposition() == "attachment":
            continue
        content_type = part.get_content_type()
        if content_type not in {"text/plain", "text/html"}:
            continue
        try:
            content = part.get_content()
        except (LookupError, UnicodeDecodeError):
            content = _decode_text(part.get_payload(decode=True) or b"")
        if not isinstance(content, str):
            continue
        if content_type == "text/plain":
            plain_parts.append(content.strip())
        else:
            html_parts.append(content)

    if plain_parts:
        return "\n\n".join(part for part in plain_parts if part)

    parser = _HTMLTextExtractor()
    for html in html_parts:
        parser.feed(html)
    return parser.text()


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _extract_csv(payload: bytes) -> str:
    sample = _decode_text(payload)
    try:
        dialect = csv.Sniffer().sniff(sample[:4096])
    except csv.Error:
        dialect = csv.excel
    rows = []
    for index, row in enumerate(csv.reader(io.StringIO(sample), dialect=dialect)):
        if index >= MAX_SPREADSHEET_ROWS:
            rows.append("[row limit reached]")
            break
        rows.append("\t".join(cell.strip() for cell in row))
    return "\n".join(rows)


def _extract_xlsx(payload: bytes) -> str:
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    chunks: list[str] = []
    row_count = 0
    try:
        for worksheet in workbook.worksheets:
            chunks.append(f"[Sheet: {worksheet.title}]")
            for row in worksheet.iter_rows(values_only=True):
                row_count += 1
                if row_count > MAX_SPREADSHEET_ROWS:
                    chunks.append("[row limit reached]")
                    return "\n".join(chunks)
                chunks.append("\t".join(_cell_text(cell) for cell in row))
    finally:
        workbook.close()
    return "\n".join(chunks)


def _extract_pdf(payload: bytes) -> str:
    reader = PdfReader(io.BytesIO(payload))
    if reader.is_encrypted:
        raise ValueError("encrypted PDFs are not supported")
    pages = []
    for index, page in enumerate(reader.pages):
        if index >= MAX_PDF_PAGES:
            pages.append("[page limit reached]")
            break
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _extract_docx(payload: bytes) -> str:
    document = Document(io.BytesIO(payload))
    chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            chunks.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(chunks)


def extract_attachment_text(filename: str, content_type: str, payload: bytes) -> str:
    """Extract text from an allowlist of formats; never execute attachments."""

    suffix = Path(filename).suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"} or content_type.startswith("text/"):
        return _extract_csv(payload) if suffix in {".csv", ".tsv"} else _decode_text(payload)
    if suffix == ".xlsx":
        return _extract_xlsx(payload)
    if suffix == ".pdf":
        return _extract_pdf(payload)
    if suffix == ".docx":
        return _extract_docx(payload)
    raise ValueError(f"unsupported attachment type: {suffix or content_type}")


def parse_eml_bytes(payload: bytes, *, source_name: str = "memory://message.eml") -> InboundEmail:
    """Parse RFC 822 bytes from Graph or another non-file email source.

    Microsoft Graph's MIME endpoint returns the same bytes stored in a local
    ``.eml`` fixture. Keeping one parser for both paths prevents production
    ingestion from behaving differently from the emails exercised in tests.
    """

    if not isinstance(payload, bytes):
        raise TypeError("email payload must be bytes")
    if not payload:
        raise ValueError("email payload is empty")

    message = BytesParser(policy=policy.default).parsebytes(payload)
    if not isinstance(message, EmailMessage):
        raise ValueError("email parser returned an unexpected message type")

    sender_name, sender_address = parseaddr(str(message.get("From", "")))
    attachment_records: list[AttachmentRecord] = []
    total_bytes = 0
    ingestion_issues: list[str] = []
    for part in message.iter_attachments():
        if len(attachment_records) >= MAX_ATTACHMENTS:
            ingestion_issues.append("attachment count exceeds extraction limit")
            break
        filename = part.get_filename() or f"attachment-{len(attachment_records) + 1}"
        payload = part.get_payload(decode=True) or b""
        total_bytes += len(payload)
        extraction_error: str | None = None
        extracted_text = ""
        if len(payload) > MAX_ATTACHMENT_BYTES:
            extraction_error = "attachment exceeds 10 MB extraction limit"
        elif total_bytes > MAX_TOTAL_ATTACHMENT_BYTES:
            extraction_error = "email exceeds 25 MB total attachment extraction limit"
        else:
            try:
                extracted_text = extract_attachment_text(
                    filename, part.get_content_type(), payload,
                )
                if len(extracted_text) > MAX_EXTRACTED_CHARACTERS:
                    extraction_error = "attachment text exceeds extraction limit"
                    extracted_text = extracted_text[:MAX_EXTRACTED_CHARACTERS]
                if any(marker in extracted_text for marker in (
                    "[row limit reached]", "[page limit reached]",
                )):
                    extraction_error = "attachment extraction was truncated"
            except Exception as exc:  # Keep a bad attachment from losing the email.
                extraction_error = str(exc)
        attachment_records.append(
            AttachmentRecord(
                filename=filename,
                content_type=part.get_content_type(),
                size_bytes=len(payload),
                sha256=hashlib.sha256(payload).hexdigest(),
                extracted_text=extracted_text,
                extraction_error=extraction_error,
            )
        )

    # Keep only headers that are operationally useful; the raw file remains the
    # audit source and may contain sensitive transport metadata.
    header_names = ("Message-ID", "Date", "From", "Reply-To", "To", "Cc", "Subject")
    headers = {name: str(message.get(name)) for name in header_names if message.get(name)}
    return InboundEmail(
        source_path=source_name,
        message_id=str(message.get("Message-ID")) if message.get("Message-ID") else None,
        sent_at=str(message.get("Date")) if message.get("Date") else None,
        sender=sender_name or sender_address,
        from_address=sender_address.lower(),
        reply_to=_addresses(message.get_all("Reply-To", [])),
        to=_addresses(message.get_all("To", [])),
        cc=_addresses(message.get_all("Cc", [])),
        subject=str(message.get("Subject", "")),
        body=_body_text(message),
        headers=headers,
        attachments=attachment_records,
        ingestion_issues=ingestion_issues,
    )


def parse_eml(path: str | Path) -> InboundEmail:
    """Parse one local ``.eml`` file with the same parser used for Graph MIME."""

    source = Path(path).resolve()
    if source.suffix.lower() != ".eml":
        raise ValueError(f"expected an .eml file, got {source.name}")
    return parse_eml_bytes(source.read_bytes(), source_name=str(source))
