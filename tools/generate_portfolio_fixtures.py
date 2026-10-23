#!/usr/bin/env python3
"""Generate local business-document fixtures backed by WWI identifiers."""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from email.message import EmailMessage
from email.policy import default
from email.utils import format_datetime
import hashlib
import json
import mimetypes
from pathlib import Path
import shutil
import sqlite3

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


DEMO_AS_OF = "2016-06-15"
NAVY = colors.HexColor("#24435C")
PALE_BLUE = colors.HexColor("#EAF1F6")
LIGHT_GRAY = colors.HexColor("#D9D9D9")


def money(value: Decimal | str | int) -> Decimal:
    return Decimal(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def write_csv(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def write_email(
    path: Path,
    *,
    sender: str,
    recipient: str,
    subject: str,
    body: str,
    sent_at: str,
    message_id: str,
    attachments: list[Path] | None = None,
    in_reply_to: str | None = None,
) -> None:
    message = EmailMessage()
    message["From"] = sender
    message["To"] = recipient
    message["Subject"] = subject
    message["Date"] = format_datetime(datetime.fromisoformat(sent_at).replace(tzinfo=timezone.utc))
    message["Message-ID"] = f"<{message_id}@portfolio.local>"
    if in_reply_to:
        message["In-Reply-To"] = f"<{in_reply_to}@portfolio.local>"
        message["References"] = f"<{in_reply_to}@portfolio.local>"
    message.set_content(body.rstrip() + "\n")
    for attachment in attachments or []:
        mime, _ = mimetypes.guess_type(attachment.name)
        maintype, subtype = (mime or "application/octet-stream").split("/", 1)
        message.add_attachment(
            attachment.read_bytes(), maintype=maintype, subtype=subtype, filename=attachment.name
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(message.as_bytes(policy=default))


def pdf_styles():
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "FixtureTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=20,
            leading=24,
            textColor=colors.black,
            alignment=TA_LEFT,
            spaceAfter=14,
        ),
        "heading": ParagraphStyle(
            "FixtureHeading",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=colors.black,
            spaceBefore=9,
            spaceAfter=6,
        ),
        "body": ParagraphStyle(
            "FixtureBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#222222"),
        ),
        "small": ParagraphStyle(
            "FixtureSmall",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#555555"),
        ),
        "table_header": ParagraphStyle(
            "FixtureTableHeader",
            parent=styles["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor=colors.white,
        ),
        "right": ParagraphStyle(
            "FixtureRight",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            alignment=TA_RIGHT,
        ),
    }


def _footer(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawString(0.62 * inch, 0.38 * inch, "Wide World Importers")
    canvas.drawRightString(7.88 * inch, 0.38 * inch, f"Page {document.page}")
    canvas.restoreState()


def build_pdf(
    path: Path,
    *,
    title: str,
    subtitle: str,
    metadata: list[tuple[str, str]],
    headers: list[str] | None = None,
    rows: list[list[str]] | None = None,
    widths: list[float] | None = None,
    totals: list[tuple[str, str]] | None = None,
    sections: list[tuple[str, str]] | None = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = pdf_styles()
    document = SimpleDocTemplate(
        str(path), pagesize=letter, rightMargin=0.62 * inch, leftMargin=0.62 * inch,
        topMargin=0.58 * inch, bottomMargin=0.55 * inch,
        title=title, author="Wide World Importers",
    )
    story = [Paragraph(title, styles["title"]), Paragraph(subtitle, styles["body"]), Spacer(1, 11)]
    meta_data = [[Paragraph(f"<b>{label}</b>", styles["small"]), Paragraph(value, styles["body"])] for label, value in metadata]
    meta = Table(meta_data, colWidths=[1.35 * inch, 5.65 * inch])
    meta.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LINEBELOW", (0, -1), (-1, -1), 0.5, LIGHT_GRAY),
    ]))
    story.extend([meta, Spacer(1, 12)])
    if headers and rows:
        table_data = [[Paragraph(h, styles["table_header"]) for h in headers]]
        table_data.extend([[Paragraph(str(cell), styles["small"]) for cell in row] for row in rows])
        table = Table(table_data, colWidths=[w * inch for w in (widths or [])], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.45, LIGHT_GRAY),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, PALE_BLUE]),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.extend([table, Spacer(1, 10)])
    if totals:
        total_data = [[Paragraph(label, styles["right"]), Paragraph(value, styles["right"])] for label, value in totals]
        total_table = Table(total_data, colWidths=[5.6 * inch, 1.4 * inch], hAlign="RIGHT")
        total_table.setStyle(TableStyle([
            ("LINEABOVE", (0, 0), (-1, 0), 0.7, NAVY),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.extend([total_table, Spacer(1, 9)])
    for heading, text in sections or []:
        story.extend([Paragraph(heading, styles["heading"]), Paragraph(text, styles["body"])])
    document.build(story, onFirstPage=_footer, onLaterPages=_footer)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def build_quote_template(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(34, 34, 34)
    title_style = document.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)

    document.add_paragraph("Customer Quote", style="Title")
    intro = document.add_paragraph()
    intro.add_run("Wide World Importers").bold = True
    intro.add_run("\nPrepared for customer review")

    details = document.add_table(rows=4, cols=4)
    details.autofit = False
    detail_rows = [
        ("Quote number", "[Quote number]", "Quote date", "[Date]"),
        ("Customer", "[Customer name]", "Valid through", "[Date]"),
        ("Contact", "[Contact name]", "Salesperson", "[Name]"),
        ("Delivery address", "[Address]", "Customer PO", "[Reference]"),
    ]
    for row, values in zip(details.rows, detail_rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[2].paragraphs[0].runs[0].bold = True
    for row in details.rows:
        row.cells[0].width = Inches(1.0)
        row.cells[1].width = Inches(2.6)
        row.cells[2].width = Inches(1.0)
        row.cells[3].width = Inches(2.6)

    document.add_paragraph()
    lines = document.add_table(rows=5, cols=6)
    headers = ["Item", "Description", "Quantity", "Unit", "Unit price", "Line total"]
    for cell, text in zip(lines.rows[0].cells, headers):
        cell.text = text
        set_cell_shading(cell, "24435C")
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    placeholders = [
        ["[SKU]", "[Product description]", "[Qty]", "[UOM]", "$[0.00]", "$[0.00]"],
        ["", "", "", "", "", ""],
        ["", "", "", "", "", ""],
        ["", "", "", "", "", ""],
    ]
    for row, values in zip(lines.rows[1:], placeholders):
        for cell, value in zip(row.cells, values):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    widths = [0.7, 2.7, 0.8, 0.7, 1.0, 1.0]
    for row in lines.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)

    totals = document.add_table(rows=4, cols=2)
    totals.alignment = 2
    for row, values in zip(totals.rows, [
        ("Subtotal", "$[0.00]"), ("Tax", "$[0.00]"),
        ("Freight", "$[0.00]"), ("Total", "$[0.00]"),
    ]):
        row.cells[0].text, row.cells[1].text = values
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    totals.rows[-1].cells[0].paragraphs[0].runs[0].bold = True
    totals.rows[-1].cells[1].paragraphs[0].runs[0].bold = True

    document.add_heading("Terms", level=1)
    document.add_paragraph(
        "Prices and availability are subject to final salesperson review. Taxes and freight are shown separately. Ambiguous product descriptions must be resolved before this quote is issued."
    )
    document.add_heading("Approval", level=1)
    document.add_paragraph("Salesperson approval  ____________________    Date  ____________")
    document.core_properties.title = "Customer Quote Template"
    document.core_properties.author = "Wide World Importers"
    document.save(path)


class FixtureBuilder:
    def __init__(self, oltp_path: Path, dw_path: Path, output: Path):
        self.output = output
        self.oltp = sqlite3.connect(f"file:{oltp_path.resolve()}?mode=ro", uri=True)
        self.oltp.row_factory = sqlite3.Row
        self.dw = sqlite3.connect(f"file:{dw_path.resolve()}?mode=ro", uri=True)
        self.dw.row_factory = sqlite3.Row

    def close(self) -> None:
        self.oltp.close()
        self.dw.close()

    def validate_spreadsheets(self) -> None:
        required = [
            "02_inbound_quote/attachments/rfq_003_store_restock.xlsx",
            "02_inbound_quote/attachments/rfq_006_mixed_units.xlsx",
            "04_supplier_delay/attachments/po_2073_split_delivery_schedule.xlsx",
        ]
        missing = [name for name in required if not (self.output / name).is_file()]
        if missing:
            raise RuntimeError(f"Generate spreadsheet fixtures first; missing: {missing}")

    def build_margin(self) -> None:
        folder = self.output / "01_margin_operations"
        pricing = self.oltp.execute(
            """
            SELECT il.invoice_line_id, il.invoice_id, i.invoice_date, i.customer_id,
                   il.stock_item_id, si.stock_item_name, si.unit_price AS catalog_price,
                   il.unit_price AS applied_price, il.quantity
            FROM sales__invoice_lines il
            JOIN sales__invoices i USING(invoice_id)
            JOIN warehouse__stock_items si USING(stock_item_id)
            WHERE i.invoice_date IN ('2016-04-21','2016-04-22')
              AND CAST(il.unit_price AS NUMERIC) < CAST(si.unit_price AS NUMERIC)
            ORDER BY i.invoice_date, il.invoice_line_id
            """
        ).fetchall()
        pricing_rows = []
        for row in pricing:
            pricing_rows.append({
                "invoice_line_id": row["invoice_line_id"],
                "invoice_id": row["invoice_id"],
                "invoice_date": row["invoice_date"],
                "customer_id": row["customer_id"],
                "stock_item_id": row["stock_item_id"],
                "stock_item_name": row["stock_item_name"],
                "catalog_unit_price": row["catalog_price"],
                "applied_unit_price": row["applied_price"],
                "quantity": row["quantity"],
                "pricing_reason": "Q2 Tailspin USB promotion",
                "source_special_deal_id": 2,
                "approved_by": "Pricing desk",
            })
        write_csv(
            folder / "inputs/applied_pricing_decisions.csv",
            list(pricing_rows[0]), pricing_rows,
        )
        cost_rows = []
        for index, row in enumerate(pricing[:6]):
            cost_rows.append({
                "invoice_line_id": row["invoice_line_id"],
                "adjustment_type": "Inbound freight allocation" if index < 5 else "Cost unavailable",
                "adjustment_amount": f"{money(Decimal('0.12') * (index + 1))}" if index < 5 else "",
                "cost_status": "RECORDED" if index < 5 else "MISSING",
                "source_reference": f"LANDED-2016-04-{index + 1:03d}" if index < 5 else "",
            })
        write_csv(
            folder / "inputs/landed_cost_adjustments.csv",
            list(cost_rows[0]), cost_rows,
        )
        comparison = self.dw.execute(
            """
            SELECT invoice_date_key, ROUND(SUM(CAST(total_excluding_tax AS REAL)),2) sales,
                   ROUND(SUM(CAST(profit AS REAL)),2) profit
            FROM fact__sale WHERE invoice_date_key IN ('2016-04-21','2016-04-22')
            GROUP BY invoice_date_key ORDER BY invoice_date_key
            """
        ).fetchall()
        write_json(folder / "expected/margin_question.json", {
            "question": "Sales increased, so why did gross profit fall?",
            "comparison_dates": [dict(row) for row in comparison],
            "required_breakdowns": ["customer", "product", "recorded selling price", "recorded cost"],
            "evidence_policy": "Promotion causality requires applied_pricing_decisions.csv. Missing cost remains unavailable.",
        })

    def build_quote(self) -> None:
        folder = self.output / "02_inbound_quote"
        attachment_3 = folder / "attachments/rfq_003_store_restock.xlsx"
        attachment_6 = folder / "attachments/rfq_006_mixed_units.xlsx"
        requests = [
            ("rfq_001_exact.eml", "waldemar@tailspintoys.com", "Quote for green USB launchers",
             "Please quote 24 USB missile launcher (Green), stock item 1. We need delivery by June 24.", None),
            ("rfq_002_ambiguous.eml", "lorena@tailspintoys.com", "USB launcher quote",
             "Can you quote 50 USB launchers? Gray or green was not specified by the store manager. Please let me know what you need.", None),
            ("rfq_003_spreadsheet.eml", "waldemar@tailspintoys.com", "Store restock spreadsheet",
             "Please quote the attached restock list for Tailspin Toys head office.", attachment_3),
            ("rfq_004_misspelling.eml", "lorena@tailspintoys.com", "Bananna flash drives",
             "We need 100 bananna USB food flash drives. Please quote delivery next week.", None),
            ("rfq_005_missing_spec.eml", "waldemar@tailspintoys.com", "Shipping cartons",
             "Please quote 20 outers of brown shipping cartons. I do not have the dimensions yet.", None),
            ("rfq_006_mixed_units.eml", "lorena@tailspintoys.com", "Packaging order mixed units",
             "The attached request mixes outers and each. Please confirm the converted quantities on the quote.", attachment_6),
            ("rfq_007_insufficient_stock.eml", "waldemar@tailspintoys.com", "Large carton requirement",
             "Please quote 10,000 outers of stock item 184 for delivery by June 20. Partial shipment is not acceptable.", None),
            ("rfq_008_customer_price.eml", "waldemar@tailspintoys.com", "Q2 Tailspin pricing",
             "Please quote 40 USB food variety packs and apply our current Q2 Tailspin agreement.", None),
        ]
        for index, (filename, sender, subject, body, attachment) in enumerate(requests, 1):
            write_email(
                folder / "emails" / filename,
                sender=sender, recipient="quotes@wideworldimporters.example",
                subject=subject, body=body, sent_at=f"2016-06-{index + 1:02d}T10:15:00",
                message_id=f"rfq-{index:03d}", attachments=[attachment] if attachment else None,
            )
        write_csv(folder / "attachments/rfq_004_followup_items.csv",
                  ["description", "quantity", "unit", "required_by"],
                  [{"description": "bananna USB food flash drive", "quantity": 100,
                    "unit": "Each", "required_by": "2016-06-24"}])
        aliases = [
            {"alias": "green USB launcher", "stock_item_id": 1, "canonical_name": "USB missile launcher (Green)"},
            {"alias": "gray USB launcher", "stock_item_id": 2, "canonical_name": "USB rocket launcher (Gray)"},
            {"alias": "bananna flash drive", "stock_item_id": 9, "canonical_name": "USB food flash drive - banana"},
            {"alias": "dessert variety pack", "stock_item_id": 15, "canonical_name": "USB food flash drive - dessert 10 drive variety pack"},
            {"alias": "305 mm brown carton", "stock_item_id": 184, "canonical_name": "Shipping carton (Brown) 305x305x305mm"},
            {"alias": "red tape dispenser", "stock_item_id": 204, "canonical_name": "Tape dispenser (Red)"},
        ]
        write_csv(folder / "reference/product_aliases.csv", list(aliases[0]), aliases)
        conversions = [
            {"stock_item_id": 184, "from_unit": "Outer", "to_unit": "Each", "multiplier": 25},
            {"stock_item_id": 191, "from_unit": "Outer", "to_unit": "Each", "multiplier": 36},
            {"stock_item_id": 204, "from_unit": "Outer", "to_unit": "Each", "multiplier": 10},
            {"stock_item_id": 15, "from_unit": "Packet", "to_unit": "Each", "multiplier": 1},
        ]
        write_csv(folder / "reference/uom_conversions.csv", list(conversions[0]), conversions)
        agreements = [
            {"agreement_id": "TT-Q2-2016", "customer_id": 1, "buying_group_id": 1,
             "stock_group_id": 7, "start_date": "2016-04-01", "end_date": "2016-06-30",
             "discount_percentage": "15.00", "requires_review": "false"},
            {"agreement_id": "TT-HO-VOLUME", "customer_id": 1, "buying_group_id": "",
             "stock_group_id": "", "start_date": "2016-06-01", "end_date": "2016-08-31",
             "discount_percentage": "5.00", "requires_review": "true"},
        ]
        write_csv(folder / "reference/customer_pricing_agreements.csv", list(agreements[0]), agreements)
        build_pdf(
            folder / "reference/tailspin_q2_pricing_agreement.pdf",
            title="Tailspin Toys Pricing Agreement",
            subtitle="Commercial pricing terms for the second quarter of 2016.",
            metadata=[("Agreement", "TT-Q2-2016"), ("Customer", "Tailspin Toys (Head Office), WWI customer 1"),
                      ("Effective", "April 1 through June 30, 2016"), ("Approval", "Pricing desk")],
            headers=["Eligible products", "Discount", "Order unit", "Freight"],
            rows=[["Novelty items in stock group 7", "15.00%", "Each or packet", "Quoted separately"]],
            widths=[3.0, 1.1, 1.4, 1.5],
            sections=[("Conditions", "The agreement applies only to eligible products ordered during the effective period. Product identity and quantity must be confirmed before the price is issued.")],
        )
        build_pdf(
            folder / "reference/customer_purchase_order_tt_160615.pdf",
            title="Customer Purchase Order",
            subtitle="Tailspin Toys purchasing request.",
            metadata=[("Customer PO", "TT-160615"), ("Customer", "Tailspin Toys (Head Office)"),
                      ("Order date", "June 15, 2016"), ("Required date", "June 24, 2016")],
            headers=["Requested item", "Quantity", "Unit"],
            rows=[["USB missile launcher (Green)", "24", "Each"],
                  ["USB food flash drive - banana", "100", "Each"]],
            widths=[5.0, 1.0, 1.0],
            sections=[("Delivery", "Deliver to the customer address on file. Confirm availability before accepting this purchase order.")],
        )
        build_quote_template(folder / "templates/draft_quote_template.docx")
        write_json(folder / "expected/outcomes.json", {
            "rfq_001": {"status": "matched", "stock_item_id": 1},
            "rfq_002": {"status": "needs_clarification", "candidate_stock_item_ids": [1, 2]},
            "rfq_003": {"status": "matched", "stock_item_ids": [9, 10, 15]},
            "rfq_004": {"status": "matched_with_alias", "stock_item_id": 9},
            "rfq_005": {"status": "needs_clarification", "reason": "carton dimensions missing"},
            "rfq_006": {"status": "matched_with_uom_conversion", "stock_item_ids": [184, 191, 204]},
            "rfq_007": {"status": "insufficient_stock", "stock_item_id": 184},
            "rfq_008": {"status": "pricing_review", "agreement_id": "TT-Q2-2016"},
        })

    def po_line(self, line_id: int) -> sqlite3.Row:
        row = self.oltp.execute(
            """
            SELECT pol.*, po.supplier_id, s.supplier_name, si.stock_item_name
            FROM purchasing__purchase_order_lines pol
            JOIN purchasing__purchase_orders po USING(purchase_order_id)
            JOIN purchasing__suppliers s USING(supplier_id)
            JOIN warehouse__stock_items si USING(stock_item_id)
            WHERE purchase_order_line_id=?
            """, (line_id,)
        ).fetchone()
        if row is None:
            raise RuntimeError(f"Missing WWI purchase-order line {line_id}")
        return row

    def build_ap(self) -> None:
        folder = self.output / "03_invoice_matching"
        scenarios = [
            {"id": "ap_001_clean", "invoice": "FAB-160601", "lines": [(8347, 2, "84.00")], "tax": "15", "freight": "0", "expected": "approve"},
            {"id": "ap_003_partial_receipt", "invoice": "FAB-160603", "lines": [(8359, 100, "84.00")], "tax": "15", "freight": "0", "expected": "exception_partial_receipt"},
            {"id": "ap_004_quantity_error", "invoice": "LIT-160604", "lines": [(8365, 1800, "47.50")], "tax": "15", "freight": "0", "expected": "exception_quantity"},
            {"id": "ap_005_price_error", "invoice": "FAB-160605", "lines": [(8350, 3385, "90.00")], "tax": "15", "freight": "0", "expected": "exception_price"},
            {"id": "ap_006_unexpected_freight", "invoice": "LIT-160606", "lines": [(8326, 1644, "47.50")], "tax": "15", "freight": "850.00", "expected": "exception_freight"},
            {"id": "ap_007_tax_error", "invoice": "LIT-160607", "lines": [(8344, 1672, "47.50")], "tax": "8", "freight": "0", "expected": "exception_tax"},
            {"id": "ap_008_mixed", "invoice": "FAB-160608", "lines": [(8338, 3377, "84.00"), (8339, 3653, "88.00"), (8340, 1822, "84.00")], "tax": "15", "freight": "0", "expected": "exception_mixed"},
        ]
        receipt_quantities = {8347: 2, 8359: 60, 8365: 1674, 8350: 3385, 8326: 1644,
                              8344: 1672, 8338: 3377, 8339: 1000, 8340: 1822}
        receipt_rows = []
        expected = []
        for index, scenario in enumerate(scenarios, 1):
            details = []
            subtotal = Decimal("0")
            supplier_name = ""
            supplier_id = 0
            po_ids = set()
            for line_id, quantity, unit_price in scenario["lines"]:
                source = self.po_line(line_id)
                supplier_name = source["supplier_name"]
                supplier_id = source["supplier_id"]
                po_ids.add(source["purchase_order_id"])
                amount = money(Decimal(quantity) * Decimal(unit_price))
                subtotal += amount
                details.append([str(source["purchase_order_id"]), str(line_id), source["stock_item_name"],
                                f"{quantity:,}", f"${Decimal(unit_price):,.2f}", f"${amount:,.2f}"])
                receipt_rows.append({
                    "receipt_id": f"GRN-{line_id}", "received_at": f"2016-06-{index + 1:02d}T08:00:00",
                    "purchase_order_id": source["purchase_order_id"], "purchase_order_line_id": line_id,
                    "stock_item_id": source["stock_item_id"], "received_outers": receipt_quantities[line_id],
                    "delivery_note": f"DN-{source['purchase_order_id']}-{line_id}",
                })
            freight = Decimal(scenario["freight"])
            tax = money((subtotal + freight) * Decimal(scenario["tax"]) / Decimal("100"))
            total = money(subtotal + freight + tax)
            pdf_path = folder / "invoices" / f"{scenario['id']}_{scenario['invoice']}.pdf"
            build_pdf(
                pdf_path,
                title="Supplier Invoice",
                subtitle=f"Invoice from {supplier_name} to Wide World Importers.",
                metadata=[("Invoice number", scenario["invoice"]), ("Invoice date", f"June {index + 1}, 2016"),
                          ("Supplier", f"{supplier_name}, WWI supplier {supplier_id}"),
                          ("Purchase order", ", ".join(map(str, sorted(po_ids))))],
                headers=["PO", "PO line", "Description", "Qty outer", "Unit price", "Amount"],
                rows=details, widths=[0.55, 0.65, 3.2, 0.75, 0.85, 1.0],
                totals=[("Subtotal", f"${subtotal:,.2f}"), ("Freight", f"${freight:,.2f}"),
                        (f"Tax {scenario['tax']}%", f"${tax:,.2f}"), ("Invoice total", f"${total:,.2f}")],
                sections=[("Payment", "Payment is due under the supplier terms on file. Quote the invoice number with remittance.")],
            )
            write_email(
                folder / "emails" / f"{scenario['id']}.eml",
                sender=f"accounts@{supplier_name.lower().replace(',','').replace(' ','')}.example",
                recipient="accounts.payable@wideworldimporters.example",
                subject=f"Invoice {scenario['invoice']}",
                body=f"Please find invoice {scenario['invoice']} attached for purchase order {', '.join(map(str, sorted(po_ids)))}.",
                sent_at=f"2016-06-{index + 1:02d}T09:20:00", message_id=scenario["id"], attachments=[pdf_path],
            )
            expected.append({"scenario_id": scenario["id"], "invoice_number": scenario["invoice"],
                             "expected_result": scenario["expected"]})

        clean = folder / "invoices/ap_001_clean_FAB-160601.pdf"
        duplicate = folder / "invoices/ap_002_duplicate_renamed.pdf"
        shutil.copyfile(clean, duplicate)
        write_email(
            folder / "emails/ap_002_duplicate.eml", sender="accounts@fabrikaminc.example",
            recipient="accounts.payable@wideworldimporters.example", subject="Invoice copy FAB-160601",
            body="Resending invoice FAB-160601 with a different filename in case the first message was missed.",
            sent_at="2016-06-03T11:10:00", message_id="ap-002-duplicate", attachments=[duplicate],
            in_reply_to="ap_001_clean",
        )
        expected.insert(1, {"scenario_id": "ap_002_duplicate", "invoice_number": "FAB-160601",
                            "expected_result": "duplicate_no_new_payable"})
        write_csv(folder / "receipts/goods_received.csv", list(receipt_rows[0]), receipt_rows)
        build_pdf(
            folder / "reference/supplier_freight_terms.pdf", title="Supplier Freight Terms",
            subtitle="Freight and surcharge terms for selected purchasing fixtures.",
            metadata=[("Effective date", "January 1, 2016"), ("Applies to", "Fabrikam, Inc. and Litware, Inc."),
                      ("Buyer", "Wide World Importers"), ("Approval reference", "PUR-TERMS-2016")],
            headers=["Charge", "Rule", "Approval required"],
            rows=[["Standard freight", "Included in agreed PO price", "No"],
                  ["Expedited freight", "Separate charge", "Written buyer approval"],
                  ["Fuel surcharge", "Not permitted", "Written contract amendment"]],
            widths=[1.5, 3.8, 1.7],
            sections=[("Invoice handling", "Accounts payable must hold any separate freight or surcharge that lacks the required approval reference.")],
        )
        write_json(folder / "expected/outcomes.json", expected)

    def build_supplier_delay(self) -> None:
        folder = self.output / "04_supplier_delay"
        messages = [
            ("supplier_001_po2073_ack.eml", "billl@fabrikam.com", "PO 2073 acknowledgment",
             "We acknowledge purchase order 2073 and confirm all six lines for delivery on June 20, 2016.", "2016-06-01T09:00:00", None),
            ("supplier_002_po2074_ack.eml", "eliasm@litwareinc.com", "PO 2074 acknowledgment",
             "Litware confirms purchase order 2074 at the ordered quantities and prices for delivery on June 20.", "2016-06-01T10:00:00", None),
            ("supplier_003_po2073_uncertain.eml", "billl@fabrikam.com", "Re PO 2073 possible delay",
             "The shirt order may slip by about a week because one production run is still being checked. This is not yet confirmed.", "2016-06-08T14:10:00", "supplier-001"),
            ("supplier_004_po2073_confirmed.eml", "billl@fabrikam.com", "Re PO 2073 revised dates",
             "The delay is now confirmed. Lines 8359 and 8360 will split between June 27 and July 5. Line 8362 has a final balance that may move to July 12. See the attached schedule.", "2016-06-10T09:35:00", "supplier-001"),
            ("supplier_005_po2073_shortfall.eml", "billl@fabrikam.com", "PO 2073 quantity shortfall",
             "We can supply only 3,400 outers on line 8362 by July 5. The remaining 368 outers are not confirmed.", "2016-06-11T16:05:00", "supplier-001"),
            ("supplier_006_po2074_price.eml", "eliasm@litwareinc.com", "PO 2074 proposed price change",
             "The unit price for line 8366 would increase from 38.40 to 41.20 per outer. Please confirm whether you accept the change. No change has been booked.", "2016-06-12T11:45:00", "supplier-002"),
            ("supplier_007_po2074_asn.eml", "eliasm@litwareinc.com", "Advance shipment notice PO 2074",
             "The first shipment against PO 2074 has been dispatched. The advance shipment notice is attached.", "2016-06-14T15:20:00", "supplier-002"),
        ]
        schedule = folder / "attachments/po_2073_split_delivery_schedule.xlsx"
        asn = folder / "attachments/po_2074_advance_shipment_notice.pdf"
        build_pdf(
            asn, title="Advance Shipment Notice", subtitle="Litware shipment against purchase order 2074.",
            metadata=[("ASN", "LIT-ASN-2074-01"), ("Purchase order", "2074"),
                      ("Ship date", "June 14, 2016"), ("Estimated arrival", "June 20, 2016")],
            headers=["PO line", "Stock item", "Quantity outer", "Shipment status"],
            rows=[["8365", "184", "900", "Dispatched"], ["8366", "193", "1,400", "Dispatched"],
                  ["8367", "204", "1,600", "Dispatched"]],
            widths=[1.0, 1.2, 1.5, 3.3],
            sections=[("Balance", "The remaining quantities are scheduled for a later shipment and are not covered by this notice.")],
        )
        for index, (filename, sender, subject, body, sent_at, reply) in enumerate(messages, 1):
            attachments = [schedule] if index == 4 else [asn] if index == 7 else None
            write_email(
                folder / "emails" / filename, sender=sender,
                recipient="purchasing@wideworldimporters.example", subject=subject, body=body,
                sent_at=sent_at, message_id=f"supplier-{index:03d}", attachments=attachments,
                in_reply_to=reply,
            )
        priorities = self.oltp.execute(
            """
            SELECT DISTINCT o.order_id, o.customer_id, ol.stock_item_id, o.expected_delivery_date
            FROM sales__orders o JOIN sales__order_lines ol USING(order_id)
            WHERE ol.stock_item_id IN (77,78,80,86,95,98,184,193,204)
              AND ol.picked_quantity < ol.quantity AND o.order_date >= '2016-05-01'
            ORDER BY o.expected_delivery_date, o.order_id LIMIT 15
            """
        ).fetchall()
        priority_rows = []
        for index, row in enumerate(priorities):
            priority_rows.append({
                "order_id": row["order_id"], "customer_id": row["customer_id"],
                "stock_item_id": row["stock_item_id"], "expected_delivery_date": row["expected_delivery_date"],
                "service_level": "Priority" if index < 5 else "Standard",
            })
        write_csv(folder / "reference/customer_order_priorities.csv", list(priority_rows[0]), priority_rows)
        write_json(folder / "expected/outcomes.json", {
            "po_2073": {"issue_count": 2, "confirmed_delivery_change": "2016-06-27",
                        "unconfirmed_final_balance_date": "2016-07-12", "update_existing_issue": True},
            "po_2074": {"issue_count": 1, "price_change_status": "awaiting_buyer_confirmation",
                        "asn": "LIT-ASN-2074-01"},
        })

    def open_ar(self) -> list[sqlite3.Row]:
        return self.oltp.execute(
            """
            SELECT ct.customer_transaction_id, ct.customer_id, ct.invoice_id,
                   ct.transaction_date, ct.outstanding_balance, c.customer_name,
                   c.payment_days, p.full_name, p.email_address
            FROM sales__customer_transactions ct
            JOIN sales__customers c USING(customer_id)
            JOIN application__people p ON p.person_id=c.primary_contact_person_id
            WHERE CAST(ct.outstanding_balance AS NUMERIC) <> 0
            ORDER BY ct.invoice_id LIMIT 12
            """
        ).fetchall()

    def build_customer_invoice_pdf(self, invoice_id: int, path: Path) -> None:
        header = self.oltp.execute(
            """
            SELECT i.invoice_id, i.invoice_date, c.customer_id, c.customer_name,
                   c.payment_days, i.customer_purchase_order_number
            FROM sales__invoices i JOIN sales__customers c USING(customer_id)
            WHERE invoice_id=?
            """, (invoice_id,)
        ).fetchone()
        lines = self.oltp.execute(
            """
            SELECT invoice_line_id, description, quantity, unit_price, tax_amount, extended_price
            FROM sales__invoice_lines WHERE invoice_id=? ORDER BY invoice_line_id
            """, (invoice_id,)
        ).fetchall()
        subtotal = sum((Decimal(row["unit_price"]) * row["quantity"] for row in lines), Decimal("0"))
        tax = sum((Decimal(row["tax_amount"]) for row in lines), Decimal("0"))
        total = sum((Decimal(row["extended_price"]) for row in lines), Decimal("0"))
        rows = [[str(row["invoice_line_id"]), row["description"], f"{row['quantity']:,}",
                 f"${Decimal(row['unit_price']):,.2f}", f"${Decimal(row['extended_price']):,.2f}"] for row in lines]
        build_pdf(
            path, title="Customer Invoice", subtitle="Wide World Importers customer invoice.",
            metadata=[("Invoice", str(invoice_id)), ("Invoice date", header["invoice_date"]),
                      ("Customer", f"{header['customer_name']}, WWI customer {header['customer_id']}"),
                      ("Customer PO", header["customer_purchase_order_number"] or "Not provided")],
            headers=["Line", "Description", "Qty", "Unit price", "Extended"], rows=rows,
            widths=[0.65, 3.85, 0.65, 0.9, 0.95],
            totals=[("Subtotal", f"${subtotal:,.2f}"), ("Tax", f"${tax:,.2f}"),
                    ("Invoice total", f"${total:,.2f}")],
            sections=[("Terms", f"Payment is due {header['payment_days']} days from the invoice date.")],
        )

    def build_receivables(self) -> None:
        folder = self.output / "05_receivables"
        records = self.open_ar()
        categories = ["routine_overdue"] * 4 + ["delivery_dispute", "pricing_dispute",
                     "missing_invoice", "missing_proof_of_delivery",
                     "promise_to_pay", "promise_to_pay", "payment_pending", "payment_pending"]
        bodies = {
            "routine_overdue": "Our records show the invoice is still open. Please send a copy and confirm the balance and due date.",
            "delivery_dispute": "We are disputing this invoice because part of the delivery was reported damaged. Please hold collection activity while this is reviewed.",
            "pricing_dispute": "The unit price on the invoice does not match the price we expected. Please review it before requesting payment.",
            "missing_invoice": "We cannot locate the invoice document. Please send another copy before following up on payment.",
            "missing_proof_of_delivery": "Please provide the signed proof of delivery. Payment is on hold until we receive it.",
            "promise_to_pay": "We confirm the balance and expect to pay it on June 20, 2016. Please note the promise on our account.",
            "payment_pending": "We sent payment by EFT on June 14, 2016. The remittance reference is included in our payment file.",
        }
        expected = []
        payment_rows = []
        for index, (record, category) in enumerate(zip(records, categories), 1):
            invoice_id = record["invoice_id"]
            write_email(
                folder / "emails" / f"ar_{index:03d}_{category}_invoice_{invoice_id}.eml",
                sender=record["email_address"], recipient="receivables@wideworldimporters.example",
                subject=f"Invoice {invoice_id} - {category.replace('_', ' ')}",
                body=bodies[category], sent_at=f"2016-06-{8 + index // 3:02d}T10:{index:02d}:00",
                message_id=f"ar-{index:03d}",
            )
            action = "draft_reminder" if category == "routine_overdue" else "suppress_generic_reminder"
            expected.append({"invoice_id": invoice_id, "customer_id": record["customer_id"],
                             "category": category, "expected_action": action})
            if category == "payment_pending":
                payment_rows.append({
                    "bank_reference": f"EFT-160614-{index:03d}", "received_date": "2016-06-14",
                    "customer_id": record["customer_id"], "invoice_id": invoice_id,
                    "amount": record["outstanding_balance"], "status": "REMITTANCE_RECEIVED",
                })
        for index, record in enumerate(records[:3], 1):
            write_email(
                folder / "prior_reminders" / f"reminder_{index:03d}_invoice_{record['invoice_id']}.eml",
                sender="receivables@wideworldimporters.example", recipient=record["email_address"],
                subject=f"Payment reminder for invoice {record['invoice_id']}",
                body=f"Invoice {record['invoice_id']} for {record['outstanding_balance']} was due on June 7, 2016. Please confirm payment status.",
                sent_at=f"2016-06-{8 + index:02d}T09:00:00", message_id=f"reminder-{index:03d}",
            )
        for record in records[:4]:
            self.build_customer_invoice_pdf(
                record["invoice_id"], folder / "invoices" / f"invoice_{record['invoice_id']}.pdf"
            )
        for record in records[4:7]:
            build_pdf(
                folder / "delivery_records" / f"proof_of_delivery_invoice_{record['invoice_id']}.pdf",
                title="Proof of Delivery", subtitle="Delivery confirmation for a Wide World Importers invoice.",
                metadata=[("Invoice", str(record["invoice_id"])), ("Customer", record["customer_name"]),
                          ("Delivered", "June 2, 2016 at 14:35"), ("Received by", record["full_name"])],
                headers=["Event", "Time", "Status"],
                rows=[["Vehicle arrived", "14:21", "Recorded"], ["Delivery completed", "14:35", "Delivered"]],
                widths=[2.6, 1.2, 3.2],
                sections=[("Confirmation", "The recipient confirmed delivery. Any damage or quantity dispute requires separate review.")],
            )
        write_csv(folder / "payments/bank_receipts.csv", list(payment_rows[0]), payment_rows)
        build_pdf(
            folder / "payments/remittance_advice.pdf", title="Customer Remittance Advice",
            subtitle="Invoice allocation supplied with the June 14 EFT receipts.",
            metadata=[("Remittance date", "June 14, 2016"), ("Payment method", "EFT"),
                      ("Recipient", "Wide World Importers"), ("Allocation count", str(len(payment_rows)))],
            headers=["Bank reference", "Customer", "Invoice", "Amount"],
            rows=[[row["bank_reference"], str(row["customer_id"]), str(row["invoice_id"]),
                   f"${Decimal(row['amount']):,.2f}"] for row in payment_rows],
            widths=[2.2, 1.1, 1.2, 2.5],
            sections=[("Allocation", "Apply each receipt only to the invoice listed above. Confirm the bank deposit before closing the receivable.")],
        )
        write_email(
            folder / "emails/ar_credit_note_request.eml", sender=records[5]["email_address"],
            recipient="receivables@wideworldimporters.example",
            subject=f"Credit note request for invoice {records[5]['invoice_id']}",
            body="Please issue a credit note for the disputed price difference. The remaining undisputed balance will be paid after the adjustment is approved.",
            sent_at="2016-06-13T15:40:00", message_id="ar-credit-note-request",
        )
        write_json(folder / "expected/outcomes.json", expected)

    def build_index_and_manifest(self) -> None:
        # artifact-tool writes inspection sidecars during workbook QA. They are
        # useful while building, but they are not part of the business fixture pack.
        for sidecar in self.output.rglob("*.inspect.ndjson"):
            sidecar.unlink()
        index = """# WWI Portfolio Document Fixtures

These local fixtures supplement the Wide World Importers SQLite databases. Every scenario uses WWI record identifiers where the source provides them. Dates after May 31, 2016 represent new business documents arriving after the sample database cutoff.

## 01 Margin Operations

Applied pricing and landed-cost CSV files support causal margin explanations and missing-cost handling.

## 02 Inbound Quote

Eight RFQ emails cover exact, abbreviated, ambiguous, incomplete, spreadsheet-based, mixed-unit, stock-shortage, and contract-pricing requests.

## 03 Invoice Matching

Eight supplier invoice scenarios cover clean matching, duplicates, partial receipt, quantity, price, freight, tax, and mixed-line exceptions.

## 04 Supplier Delay

Seven messages update purchase orders 2073 and 2074 over time, including uncertain and confirmed changes.

## 05 Receivables

Twelve customer cases include routine overdue balances, disputes, missing documents, promises to pay, and payments awaiting allocation.
"""
        (self.output / "INDEX.md").write_text(index, encoding="utf-8")
        artifacts = []
        for path in sorted(self.output.rglob("*")):
            if path.is_file() and path.name != "manifest.json":
                artifacts.append({
                    "path": path.relative_to(self.output).as_posix(),
                    "bytes": path.stat().st_size,
                    "sha256": sha256(path),
                })
        manifest = {
            "fixture_version": 1,
            "generated_at_utc": datetime.now(timezone.utc).isoformat(),
            "demo_as_of_date": DEMO_AS_OF,
            "source_databases": ["wide_world_importers.sqlite3", "wide_world_importers_dw.sqlite3"],
            "artifact_count": len(artifacts),
            "artifacts": artifacts,
        }
        write_json(self.output / "manifest.json", manifest)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--oltp", type=Path, default=Path("data/sqlite/wide_world_importers.sqlite3"))
    parser.add_argument("--dw", type=Path, default=Path("data/sqlite/wide_world_importers_dw.sqlite3"))
    parser.add_argument("--output-dir", type=Path, default=Path("data/portfolio_artifacts"))
    args = parser.parse_args()
    if not args.oltp.is_file() or not args.dw.is_file():
        raise SystemExit("Build the two WWI SQLite databases before generating fixtures")
    builder = FixtureBuilder(args.oltp, args.dw, args.output_dir)
    try:
        builder.validate_spreadsheets()
        builder.build_margin()
        builder.build_quote()
        builder.build_ap()
        builder.build_supplier_delay()
        builder.build_receivables()
        builder.build_index_and_manifest()
    finally:
        builder.close()
    print(f"Generated portfolio fixtures in {args.output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
